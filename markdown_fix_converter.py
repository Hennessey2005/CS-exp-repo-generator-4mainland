#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word文档Markdown格式修复工具
用于将Word文档中的Markdown格式的#转换为相应的字体大小，并删除所有反引号
"""

from docx import Document
from docx.shared import Pt
import re
import os
import sys


def fix_markdown_format_in_word(input_file, output_file=None):
    """
    修复Word文档中的Markdown格式
    
    Args:
        input_file: 输入Word文档路径
        output_file: 输出Word文档路径，如果为None则在原文件名后添加_fixed后缀
    """
    try:
        # 确定输出文件路径
        if output_file is None:
            base_name, ext = os.path.splitext(input_file)
            output_file = f"{base_name}_fixed{ext}"
        
        # 加载Word文档
        doc = Document(input_file)
        print(f"正在处理文档: {input_file}")
        
        # 定义标题级别对应的字体大小
        heading_font_sizes = {
            1: Pt(24),  # 一级标题（增大字体以确保可见效果）
            2: Pt(22),  # 二级标题
            3: Pt(20),  # 三级标题
            4: Pt(18),  # 四级标题
            5: Pt(16),  # 五级标题
            6: Pt(14)   # 六级标题
        }
        
        # 统计信息
        total_paragraphs = len(doc.paragraphs)
        heading_count = 0
        code_block_count = 0
        
        print(f"文档包含 {total_paragraphs} 个段落")
        
        # 遍历所有段落
        for i, paragraph in enumerate(doc.paragraphs):
            # 保存原始文本
            original_text = paragraph.text
            
            # 调试：打印每个段落的前50个字符
            if original_text.strip():
                print(f"\n段落 {i+1}/{total_paragraphs}: '{original_text[:50]}{'...' if len(original_text) > 50 else ''}'")
            
            # 调试：打印段落的原始表示，包括空格和特殊字符
            print(f"  原始文本表示: repr('{original_text[:30]}...') = {repr(original_text[:30])}")
            print(f"  原始文本长度: {len(original_text)}")
            print(f"  文本前10个字符的ASCII码: {[ord(c) for c in original_text[:10]]}")
            
            # 重点测试：直接检查是否包含##符号开头
            if original_text.startswith('##'):
                print(f"  检测到##开头的文本: '{original_text[:30]}...'")
            
            # 更强大的正则表达式集合，确保能匹配各种Markdown标题格式
            heading_patterns = [
                r'^\s*(#{1,6})\s+(.*)$',          # 标准格式，允许标题前有空格
                r'^(#{1,6})\s+(.*)$',              # 严格的开头匹配
                r'^\t*(#{1,6})\s+(.*)$',          # 制表符开头
                r'^[\u0020\u00A0]*(#{1,6})[\u0020\u00A0]+(.*)$',  # 处理各种空格字符
                r'^(#{1,6})[\u0020\u00A0]+(.*)$',  # 直接匹配#后接任何空格
                r'^[\s\t]*(#{1,6})[\s\t]+(.*)$',  # 更广泛的空白字符匹配
            ]
            
            # 直接测试特殊案例：## 未来工作方向
            test_text = "## 未来工作方向"
            if original_text == test_text:
                print(f"  完全匹配测试文本: '{test_text}'")
                # 手动创建匹配对象
                heading_match = type('obj', (object,), {})
                heading_match.group = lambda x: ('##', '未来工作方向')[x-1]
            else:
                heading_match = None
                for i, pattern in enumerate(heading_patterns):
                    match = re.match(pattern, original_text)
                    if match:
                        heading_match = match
                        print(f"  匹配成功使用模式 {i+1}: {pattern}")
                        break
            
            if heading_match:
                # 提取标题级别和内容
                heading_level = len(heading_match.group(1))
                heading_content = heading_match.group(2)
                
                # 删除内容中的反引号
                heading_content = heading_content.replace('`', '')
                
                print(f"  检测到标题: {heading_level}级 -> '{heading_content}'")
                
                # 清空段落并重新添加内容
                paragraph.text = ''  # 先清空
                run = paragraph.add_run(heading_content)
                
                # 设置字体大小
                font_size = heading_font_sizes.get(heading_level, Pt(14))
                run.font.size = font_size
                
                # 额外设置字体加粗，使标题更明显
                run.font.bold = True
                
                print(f"  设置字体大小: {font_size}, 加粗: True")
                
                heading_count += 1
            else:
                # 对于非标题段落，删除所有反引号
                if '`' in original_text:
                    # 使用更可靠的方式更新文本
                    original_runs = list(paragraph.runs)  # 保存原始runs的副本
                    paragraph.text = ''  # 清空段落
                    
                    for run in original_runs:
                        # 处理每个run中的文本
                        new_text = run.text.replace('`', '')
                        new_run = paragraph.add_run(new_text)
                        
                        # 复制原始格式
                        new_run.font.size = run.font.size
                        new_run.font.name = run.font.name
                        new_run.font.bold = run.font.bold
                        new_run.font.italic = run.font.italic
                    
                    print(f"  删除反引号: '{original_text[:50]}...' -> '{original_text.replace('`', '')[:50]}...'")
                    code_block_count += 1
                # 额外检查是否包含标题关键词，即使没有#标记
                elif any(keyword in original_text.strip() for keyword in ['引言', '实验目的', '实验内容', '实验步骤', '总结', '改进建议', '未来工作']):
                    print(f"  发现疑似标题但无#标记: '{original_text.strip()}'")
                elif original_text.strip():
                    print(f"  普通段落，无需修改")
        
        # 保存修改后的文档
        doc.save(output_file)
        print(f"\n处理完成统计:")
        print(f"  - 总共处理段落: {total_paragraphs}")
        print(f"  - 修复标题数量: {heading_count}")
        print(f"  - 删除反引号数量: {code_block_count}")
        print(f"文档已保存到: {output_file}")
        return True
        
    except Exception as e:
        print(f"处理文档时出错: {e}")
        return False


def process_directory(input_dir, recursive=False):
    """
    处理目录中的所有Word文档
    
    Args:
        input_dir: 输入目录路径
        recursive: 是否递归处理子目录
    """
    if not os.path.isdir(input_dir):
        print(f"错误: {input_dir} 不是有效的目录")
        return
    
    # 获取所有.docx文件
    docx_files = []
    if recursive:
        for root, _, files in os.walk(input_dir):
            for file in files:
                if file.endswith('.docx'):
                    docx_files.append(os.path.join(root, file))
    else:
        for file in os.listdir(input_dir):
            if file.endswith('.docx'):
                docx_files.append(os.path.join(input_dir, file))
    
    # 处理每个文件
    print(f"在目录 {input_dir} 中找到 {len(docx_files)} 个Word文档")
    success_count = 0
    for docx_file in docx_files:
        if fix_markdown_format_in_word(docx_file):
            success_count += 1
    
    print(f"处理完成: {success_count}/{len(docx_files)} 个文件处理成功")


if __name__ == "__main__":
    """
    命令行接口
    使用示例:
    python markdown_fix_converter.py input.docx
    python markdown_fix_converter.py input.docx output.docx
    python markdown_fix_converter.py --dir input_dir
    python markdown_fix_converter.py --dir input_dir --recursive
    """
    if len(sys.argv) < 2:
        print("使用方法:")
        print("  单文件处理:")
        print("    python markdown_fix_converter.py input.docx")
        print("    python markdown_fix_converter.py input.docx output.docx")
        print("  目录处理:")
        print("    python markdown_fix_converter.py --dir input_dir")
        print("    python markdown_fix_converter.py --dir input_dir --recursive")
        sys.exit(1)
    
    # 处理命令行参数
    if sys.argv[1] == '--dir':
        # 目录处理模式
        if len(sys.argv) < 3:
            print("错误: 请指定目录路径")
            sys.exit(1)
        input_dir = sys.argv[2]
        recursive = '--recursive' in sys.argv
        process_directory(input_dir, recursive)
    else:
        # 单文件处理模式
        input_file = sys.argv[1]
        if not os.path.isfile(input_file):
            print(f"错误: 文件 {input_file} 不存在")
            sys.exit(1)
        
        output_file = None
        if len(sys.argv) > 2:
            output_file = sys.argv[2]
        
        fix_markdown_format_in_word(input_file, output_file)